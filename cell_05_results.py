# [v2.0 | cell 05] Results: Complai_<deal>_<date>_<version>.xlsx (the sheet), .csv (same columns) and _report.md, written into RUN_DIR.
# Requires cells 01, 02 and 04 (or 05a). No numbering columns: FRT row numbers differ between workbooks; the question id stays internal.
_PROFILE, _CITED = DEAL_PROFILE, CRITERIA_CITED_FROM_DOCS   # from cell 02b

def risk_score(ans, layer, conf):
    """Risk score 1-5 from answer x layer x confidence, computed in code: 1-3 Permissible, 4-5 Recommend for Review, N/A none."""
    l = (layer or "").lower()
    if ans == "N/A":
        return None
    if ans == "Yes":
        if "explicit" in l:
            return 1
        return 2 if conf == "High" else 3
    if ans == "Yes - deviated from standard approach":
        return 3
    if ans == "TBD":
        return 4
    if ans == "No":
        return 4 if "authoriz" in l else 5
    return 4

def risk_verdict(score):
    return "N/A - out of scope" if score is None else ("Permissible" if score <= 3 else "Recommend for Review")

rows = []
for qid in [k for k in RUN_QIDS if k in RESULTS]:
    r = RESULTS[qid]
    _rs = risk_score(r["ANSWER_norm"], r["LAYER"], r["CONFIDENCE"])
    rows.append({"Question": r["question"], "Current File Review Relevance": "Yes",
                 "Application Permissible?": r["ANSWER_norm"], "Confidence": r["CONFIDENCE"],
                 "Rationale": r["RATIONALE"], "Criteria basis": r["CRITERIA_BASIS"],
                 "Evidence basis": r["EVIDENCE_BASIS"], "Evidence retrieved": r["evidence"],
                 "Reviewer follow-up": r["REVIEWER_FOLLOW_UP"], "Paragraph References": r["ref"],
                 "_layer": r["LAYER"], "_risk": _rs if _rs is not None else "", "_verdict": risk_verdict(_rs), "_qid": qid})   # xlsx-only (hidden columns)
CSV_COLUMNS = [c for c in rows[0] if not c.startswith("_")] if rows else []

# csv = the first ten columns; the last three are hidden in the xlsx
XLSX_HEADERS = ["Question", "Current File Review Relevance", "Application Permissible?", "Confidence", "Rationale",
                "Criteria basis", "Evidence basis", "Evidence retrieved", "Reviewer follow-up", "Paragraph References",
                "Layer (cross-sector view)", "Risk score (1-5, cross-sector view)", "Permissibility verdict"]
# header comments + Legend sheet
COLUMN_GUIDE = {
 "Question": "from FRT (wording as in the mapping workbook; refinements to the template wording are shown in red there)",
 "Current File Review Relevance": "from FRT",
 "Application Permissible?": ("DRAFT answer in the FRT's own vocabulary.\n"
    "Yes = the file applied the step in line with the criteria's standard treatment.\n"
    "Yes - deviated from standard approach = the criteria name a standard/typical treatment (a benchmark value, a table "
    "assumption, a default approach) and the file used a different one within the discretion the criteria allow; the "
    "rationale quotes the standard, what the file did and the file's documented reasoning so YOU can judge whether the "
    "deviation is well reasoned.\n"
    "No = the file's treatment is not in line with the criteria, or uses a step the criteria never authorize.\n"
    "N/A = the question's precondition is not met for this transaction (e.g. a floating-rate question on a fixed-rate "
    "deal); the deciding fact is quoted.\n"
    "TBD = the condition is met (or there is none) but the file does not show how the step was applied: a documentation "
    "gap for a human to resolve, NOT a finding."),
 "Confidence": ("The tool's self-assessed reliability of the row.\n"
    "High = the deciding fact, number or step is quoted verbatim from the file and the criteria provision reads one way only.\n"
    "Medium = support is indirect or partial (step asserted but not shown, value inferred, criteria wording open to more "
    "than one reading, or condition decided from the deal profile alone).\n"
    "Low = rests on absence of evidence, an unverified characteristic or an assumption; every TBD is Low.\n"
    "Treat Low as 'verify first'."),
 "Rationale": ("3-6 sentences linking the quoted evidence to the criteria text, drafted for the FRT's rationale "
    "column. For 'Yes - deviated' rows: the standard treatment, what the file did instead, and the file's reasoning."),
 "Criteria basis": "The criteria sentence(s) the answer relies on, quoted.",
 "Evidence basis": "The passage(s) of the file the answer relies on, quoted, with document name and section.",
 "Evidence retrieved": ("Everything the tool retrieved from the file for this question, verbatim: the "
    "complete audit trail behind the row (long by design). Each reproduced section is tagged 'addresses: <criteria point>' "
    "and the closing 'NOT FOUND: ...' names the criteria points of the question's search guide with no retrieved passage "
    "(the search guide is printed per question in the _report.md file). 'Not found' means not retrieved for this "
    "question, not proof that it is absent from the file."),
 "Reviewer follow-up": "What the tool could not verify and a human must check (model file, legal documents, missing sections); 'None' when nothing is outstanding.",
 "Paragraph References": "Criteria section the question is derived from (the FRT's paragraph-reference column).",
 "Layer (cross-sector view)": ("Which kind of criteria authority the answer rests on.\n"
    "Explicit requirement = a prescribed value, formula, floor, cap, threshold or required step.\n"
    "Discretion within criteria = a range or judgment the criteria allow ('may', 'typically', 'generally', 'up to').\n"
    "Not authorized by criteria = a step the criteria never mention or authorize.\n"
    "Not applicable = condition not met.  TBD = not evidenced."),
 "Risk score (1-5, cross-sector view)": ("Deterministic crosswalk to the cross-sector 1-5 scale, computed in code from answer x "
    "layer x confidence (never assigned by the model): Yes+Explicit=1; Yes+Discretion=2 (High confidence) or 3; "
    "Yes-deviated=3; TBD=4; No=4 (not authorized) or 5 (explicit breach); N/A=no score."),
 "Permissibility verdict": "Risk score <=3 = Permissible; >=4 = Recommend for Review; N/A = out of scope.",
}

report_path = os.path.join(RUN_DIR, f"{OUTPUT_STEM}_report.md")
with open(report_path, "w", encoding="utf-8") as f:
    f.write(f"# Gemma4 draft answers — {DEAL_NAME} — run {run_stamp} — {CODE_VERSION}\n")
    f.write(f"\nMapping: {MAPPING_PATH} (criteria version tested: {CRITERIA_VERSION_LABEL or 'n/a'})\n"
            f"Documents: {[doc_name(d) for d in DEAL_DOCS]}\n")
    if _PROFILE:
        f.write(f"\n## Deal profile (injected into every Step B)\n\n```\n{_PROFILE}\n```\n")
    for n, qid in enumerate([k for k in RUN_QIDS if k in RESULTS], start=1):
        r = RESULTS[qid]
        f.write(f"\n---\n\n## {n}. ANSWER: {r['ANSWER_norm']} | Confidence: {r['CONFIDENCE']}\n")
        f.write(f"\n**Question:** {r['question']}\n\n**Criteria reference:** {r['ref']}\n")
        if qid in QBYID:
            f.write(f"\n**Search guide (the criteria points that NOT FOUND refers to):**\n\n{QBYID[qid]['summary']}\n")
        f.write(f"\n**Rationale:** {r['RATIONALE']}\n")
        f.write(f"\n**Criteria basis:** {r['CRITERIA_BASIS']}\n\n**Evidence basis:** {r['EVIDENCE_BASIS']}\n")
        f.write(f"\n**Reviewer follow-up:** {r['REVIEWER_FOLLOW_UP']}\n")
        f.write(f"\n### Evidence retrieved (verbatim)\n\n```\n{r['evidence']}\n```\n")
print("report ->", report_path)

try:
    import pandas as pd
    df = pd.DataFrame(rows)[CSV_COLUMNS]
    csv_path = os.path.join(RUN_DIR, f"{OUTPUT_STEM}.csv"); df.to_csv(csv_path, index=False)
    print("csv (same columns as the xlsx: answer + rationale + criteria/evidence basis + full retrieved evidence per question) ->", csv_path)
    _prev = df[["Application Permissible?", "Confidence", "Rationale", "Reviewer follow-up"]].astype(str)
    _prev.insert(0, "id", [r["_qid"] for r in rows])
    display(_prev)   # console preview with the internal id
except Exception as e:
    print("csv/display step problem:", e)
    for r in rows: print(r["_qid"], r["Application Permissible?"], r["Confidence"], "|", r["Rationale"][:120])
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Gemma_draft_answers"
    ws.append(XLSX_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        if c.value in COLUMN_GUIDE:
            c.comment = Comment(COLUMN_GUIDE[c.value], "File Review AI agent"); c.comment.width = 460; c.comment.height = 230
    NON_AI_HEADERS = {"Question", "Current File Review Relevance", "Paragraph References"}   # grey fill = not model output
    NON_AI_FILL = PatternFill("solid", fgColor="EDEDED")
    _non_ai_cols = [i for i, h in enumerate(XLSX_HEADERS, start=1) if h in NON_AI_HEADERS]
    _hidden_cols = list(range(XLSX_HEADERS.index("Paragraph References") + 2, len(XLSX_HEADERS) + 1))
    for r in rows:
        ws.append([r[h] for h in XLSX_HEADERS[:10]] + [r["_layer"], r["_risk"], r["_verdict"]])
    for col, w in zip("ABCDEFGHIJKLM", [55, 12, 14, 10, 55, 38, 38, 80, 38, 28, 22, 10, 20]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in _non_ai_cols:
        for c in ws[get_column_letter(i)]: c.fill = NON_AI_FILL
    for i in _hidden_cols:
        ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.freeze_panes = "B2"   # question stays visible
    lg = wb.create_sheet("Legend"); lg.append(["Column", "Meaning and possible values"])
    for c in lg[1]: c.font = Font(bold=True)
    for h in XLSX_HEADERS: lg.append([h, COLUMN_GUIDE[h]])
    lg.append(["", ""]); lg.append(["How to read a row", "Every answer is a DRAFT. 'Evidence basis' and 'Evidence retrieved' "
        "let you verify each claim against the file; 'Reviewer follow-up' lists what the tool could not verify. "
        "TBD and Low-confidence rows are the ones to check first; N/A rows quote the fact that takes the question out of scope. "
        "Rows carry no FRT row number (numbering differs between FRT workbooks): match a row to the FRT on the question text."])
    lg.append(["Shading / hidden columns", "Grey-filled columns are NOT generated by the model (the question and its criteria "
        "reference from the mapping workbook, the template's relevance flag); unfilled columns are the model's output. "
        "The columns after 'Paragraph References' (Layer, risk score, verdict = the cross-sector view) are hidden by "
        "default — unhide them in Excel if you want that view."])
    lg.append(["Run", f"{DEAL_NAME} | run {run_stamp} | code {CODE_VERSION} | mapping {os.path.basename(MAPPING_PATH)} "
        f"[{CRITERIA_VERSION_LABEL or 'n/a'}]"])
    lg.column_dimensions["A"].width = 34; lg.column_dimensions["B"].width = 120
    for row in lg.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    xlsx_path = os.path.join(RUN_DIR, f"{OUTPUT_STEM}.xlsx"); wb.save(xlsx_path); print("xlsx ->", xlsx_path)
except Exception as e:
    print("openpyxl step skipped:", e)
